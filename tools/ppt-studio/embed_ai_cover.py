#!/usr/bin/env python3
"""Put the AI cover behind native title text."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path("/Users/imooc/Desktop/recording-kg")
IMG = ROOT / "recordings" / "ppt-studio" / "ai-cover.jpg"
PPTX = ROOT / "recordings" / "ppt-studio" / "ai-cover.pptx"


def main() -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(str(IMG), Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    box = slide.shapes.add_textbox(Inches(0.8), Inches(4.8), Inches(11.5), Inches(2.0))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = "现在完成时"
    run.font.size = Pt(48)
    run.font.bold = True
    run.font.color.rgb = RGBColor(255, 255, 255)
    run.font.name = "Heiti SC"
    p2 = tf.add_paragraph()
    run2 = p2.add_run()
    run2.text = "文生图作背景 · 标题是可改的字"
    run2.font.size = Pt(20)
    run2.font.color.rgb = RGBColor(255, 255, 255)
    run2.font.name = "Heiti SC"
    prs.save(PPTX)
    print(PPTX)


if __name__ == "__main__":
    main()
