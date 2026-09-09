#!/usr/bin/env python3
"""Demo: embed a still image and an MP4 into a native PPTX."""

from __future__ import annotations

import subprocess
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.opc.constants import CONTENT_TYPE as CT
from pptx.util import Inches, Pt

ROOT = Path("/Users/imooc/Desktop/recording-kg")
VIDEO = ROOT / "recordings" / "demo-window-only" / "lesson.mp4"
OUT_DIR = ROOT / "recordings" / "ppt-studio"
STILL = OUT_DIR / "still.png"
POSTER = OUT_DIR / "video-poster.png"
PPTX = OUT_DIR / "embed-media.pptx"
FONT = Path("/System/Library/Fonts/STHeiti Medium.ttc")


def font(size: int) -> ImageFont.FreeTypeFont:
    if FONT.exists():
        return ImageFont.truetype(str(FONT), size=size, index=0)
    return ImageFont.load_default()


def make_still() -> None:
    img = Image.new("RGB", (1600, 900), (251, 247, 240))
    draw = ImageDraw.Draw(img)
    draw.rectangle([0, 0, 1600, 18], fill=(180, 35, 24))
    draw.text((80, 280), "随便找的一张图", font=font(72), fill=(27, 20, 12))
    draw.text((80, 420), "生成后嵌进 PPT 第二页", font=font(40), fill=(109, 94, 72))
    draw.text((80, 720), "recording-kg · 嵌入验证", font=font(28), fill=(109, 94, 72))
    img.save(STILL, "PNG")


def grab_poster() -> None:
    subprocess.run(
        [
            "ffmpeg", "-y", "-ss", "3", "-i", str(VIDEO),
            "-frames:v", "1", "-q:v", "2", str(POSTER),
        ],
        check=True,
        capture_output=True,
    )


def textbox(slide, left, top, width, height, text, size, color, bold=False):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.name = "Heiti SC"


def build() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    make_still()
    grab_poster()

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    ink = RGBColor(0x1B, 0x14, 0x0C)
    muted = RGBColor(0x6D, 0x5E, 0x48)
    red = RGBColor(0xB4, 0x23, 0x18)

    s1 = prs.slides.add_slide(blank)
    bar = s1.shapes.add_shape(1, Inches(0), Inches(0), Inches(13.333), Inches(0.12))
    bar.fill.solid()
    bar.fill.fore_color.rgb = red
    bar.line.fill.background()
    textbox(s1, Inches(0.9), Inches(2.3), Inches(11.5), Inches(1.4), "图 + 视频嵌进 PPT", 44, ink, True)
    textbox(s1, Inches(0.9), Inches(4.0), Inches(11.5), Inches(1.2), "第二页是生成的配图，第三页是窗口录屏，点一下就能播。", 22, muted)

    s2 = prs.slides.add_slide(blank)
    textbox(s2, Inches(0.7), Inches(0.25), Inches(12), Inches(0.5), "配图页", 18, muted)
    s2.shapes.add_picture(str(STILL), Inches(0.7), Inches(0.85), Inches(12.0), Inches(6.35))

    s3 = prs.slides.add_slide(blank)
    textbox(s3, Inches(0.7), Inches(0.25), Inches(12), Inches(0.5), "视频页 · 单击播放", 18, muted)
    s3.shapes.add_movie(
        str(VIDEO),
        Inches(0.7),
        Inches(0.85),
        Inches(12.0),
        Inches(6.35),
        poster_frame_image=str(POSTER),
        mime_type=CT.MP4,
    )

    prs.save(PPTX)
    print(PPTX)
    print("still", STILL.stat().st_size)
    print("poster", POSTER.stat().st_size)
    print("pptx", PPTX.stat().st_size)


if __name__ == "__main__":
    build()
